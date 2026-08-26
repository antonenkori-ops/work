"""Экспорт релиза в .xlsx, визуально похожий на наброски/План работ.xlsx:
жирный заголовок сверху, разделы жирным без заливки, шапка таблицы с цветной
заливкой, строки-маркеры выделены, таблица рисков со своей (оливковой) шапкой.
"""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.models import ReleaseItem, ReleasePlan

SECTION_LABELS = {
    "prep": "Подготовительные работы",
    "main": "Основные работы",
    "closing": "Заключительные работы",
    "rollback": "План отката",
}
SECTION_ORDER = ["prep", "main", "closing", "rollback"]

# Столбцы повторяют наброски/План работ.xlsx: план (Дата и время начала/окончания)
# заполняется этим инструментом, а Начало/Конец/Время — колонки для второго
# администратора, куда он вручную вписывает факт по ходу выполнения работ.
COLUMN_WIDTHS = [5.5, 60, 16, 14, 16, 16, 10, 10, 10, 35, 55]
WORK_HEADERS = [
    "№",
    "Работы",
    "Зависит от №",
    "Продолжительность",
    "Дата и время начала",
    "Дата и время окончания",
    "Начало",
    "Конец",
    "Время",
    "Отв. Исполнитель",
    "Комментарий",
]
NUM_COLUMNS = len(COLUMN_WIDTHS)

# Цвета взяты из исходного файла: шапка таблицы — тема Office "accent3"
# (9BBB59) с tint -0.25, шапка таблицы рисков — прямой RGB 76923C.
HEADER_FILL = PatternFill("solid", fgColor="FF748C43")
MARKER_FILL = PatternFill("solid", fgColor="FFE8E8E8")
RISK_HEADER_FILL = PatternFill("solid", fgColor="FF76923C")
THIN = Side(style="thin", color="FF808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TITLE_FONT = Font(bold=True, size=16)
SECTION_FONT = Font(bold=True, size=14)
HEADER_FONT = Font(bold=True, size=11, color="FFFFFFFF")
BODY_FONT = Font(size=11)
TOTAL_FONT = Font(bold=True, size=11)
TOTAL_FILL = PatternFill("solid", fgColor="FFF3F4F6")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _fmt_duration(minutes: int | None) -> str:
    if minutes is None:
        return ""
    hours, mins = divmod(int(minutes), 60)
    return f"{hours}:{mins:02d}"


def _fmt_dt(value: datetime | None) -> str:
    return value.strftime("%d.%m.%Y %H:%M") if value else ""


def _write_merged(ws: Worksheet, row: int, text: str, font: Font, fill: PatternFill | None = None) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMN_WIDTHS))
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = font
    cell.alignment = WRAP_CENTER
    if fill is not None:
        for col in range(1, len(COLUMN_WIDTHS) + 1):
            ws.cell(row=row, column=col).fill = fill


def _write_header_row(ws: Worksheet, row: int, headers: list[str], fill: PatternFill) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = fill
        cell.alignment = WRAP_CENTER
        cell.border = BORDER


def _item_title_with_modules(item: ReleaseItem, sets_by_id: dict) -> str:
    text = getattr(item, "title_display", None) or item.title or ""
    module_set = sets_by_id.get(item.module_set_id) if item.module_set_id else None
    if module_set and module_set.entries:
        lines = "\n".join(
            f"{e.name}:{e.version}" if e.version else e.name for e in module_set.entries
        )
        text = f"{text}\n\nМодули ({module_set.name}):\n{lines}"
    return text


def build_workbook(release: ReleasePlan) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "План работ"
    ws.sheet_view.showGridLines = False

    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[chr(64 + col)].width = width

    row = 1
    title_lines = [f"План работ по внедрению релиза {release.title}"]
    for ticket in sorted(release.tickets, key=lambda t: t.sort_order):
        title_lines.append(f"{ticket.label} ({ticket.key})")
    _write_merged(ws, row, "\n".join(title_lines), TITLE_FONT)
    ws.row_dimensions[row].height = max(30, 16 * len(title_lines))
    row += 2

    marker_letters = iter("ABCDEFGHIJKLMNOPQRSTUVWXYZ")
    sets_by_id = {s.id: s for s in release.module_sets}
    items_by_section: dict[str, list[ReleaseItem]] = {s: [] for s in SECTION_ORDER}
    for item in release.items:
        items_by_section.setdefault(item.section, []).append(item)

    for section in SECTION_ORDER:
        items = sorted(items_by_section.get(section, []), key=lambda i: i.sort_order)

        _write_merged(ws, row, SECTION_LABELS[section], SECTION_FONT)
        row += 1

        _write_header_row(ws, row, WORK_HEADERS, HEADER_FILL)
        row += 1

        number_by_id = {}
        n = 0
        for item in items:
            if item.item_type == "work":
                n += 1
                number_by_id[item.id] = n

        for item in items:
            if item.item_type == "marker":
                letter = next(marker_letters, "*")
                text = f"{letter} — {_fmt_dt(item.marker_at)}\n{item.title}".strip()
                _write_merged(ws, row, text, Font(bold=True, size=12), MARKER_FILL)
                row += 1
                continue

            depends_number = number_by_id.get(item.depends_on_id, "") if item.depends_on_id else ""
            values = [
                number_by_id.get(item.id, ""),
                _item_title_with_modules(item, sets_by_id),
                depends_number,
                _fmt_duration(item.duration_minutes),
                _fmt_dt(item.start_at),
                _fmt_dt(item.end_at),
                "",  # Начало (факт) — заполняется вручную вторым администратором
                "",  # Конец (факт)
                "",  # Время (факт)
                getattr(item, "executor_display", None) or item.executor or "",
                getattr(item, "comment_display", None) or item.comment or "",
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.alignment = WRAP_CENTER if col in (1, 3, 4, 5, 6, 7, 8, 9) else WRAP_TOP
            row += 1

        work_minutes = sum(i.duration_minutes or 0 for i in items if i.item_type == "work")
        if any(i.item_type == "work" for i in items):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            label_cell = ws.cell(row=row, column=1, value="Итого по разделу")
            label_cell.font = TOTAL_FONT
            label_cell.alignment = WRAP_CENTER
            dur_cell = ws.cell(row=row, column=4, value=_fmt_duration(work_minutes))
            dur_cell.font = TOTAL_FONT
            dur_cell.alignment = WRAP_CENTER
            for col in range(1, NUM_COLUMNS + 1):
                ws.cell(row=row, column=col).border = BORDER
                ws.cell(row=row, column=col).fill = TOTAL_FILL
            row += 1

        if section == "rollback" and release.rollback_note:
            _write_merged(ws, row, release.rollback_note, Font(italic=True, size=11))
            row += 1

        row += 1

    # Таблица рисков — своя раскладка колонок в той же сетке из 11 столбцов:
    # № | Описание риска (2-4) | Уровень риска | Компенсирующие меры (6-8) | ФИО (9-11)
    _write_merged(ws, row, "Риски при внедрении", SECTION_FONT)
    row += 1

    header_row = row
    ws.merge_cells(start_row=header_row, start_column=2, end_row=header_row, end_column=4)
    ws.merge_cells(start_row=header_row, start_column=6, end_row=header_row, end_column=8)
    ws.merge_cells(start_row=header_row, start_column=9, end_row=header_row, end_column=11)
    for col, text in [
        (1, "№"),
        (2, "Описание риска"),
        (5, "Уровень риска"),
        (6, "Компенсирующие меры"),
        (9, "ФИО"),
    ]:
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.alignment = WRAP_CENTER
    for col in range(1, NUM_COLUMNS + 1):
        ws.cell(row=header_row, column=col).fill = RISK_HEADER_FILL
        ws.cell(row=header_row, column=col).border = BORDER
    row += 1

    for risk in sorted(release.risks, key=lambda r: r.sort_order):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=8)
        ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=11)
        for col, value in [
            (1, getattr(risk, "number", None)),
            (2, risk.description or ""),
            (5, risk.level or ""),
            (6, risk.measures or ""),
            (9, risk.owners or ""),
        ]:
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP_TOP if col in (2, 6, 9) else WRAP_CENTER
        for col in range(1, NUM_COLUMNS + 1):
            ws.cell(row=row, column=col).border = BORDER
        row += 1

    return wb
